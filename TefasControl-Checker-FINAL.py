
import openpyxl
import asyncio
from playwright.async_api import async_playwright

INPUT_FILE = "TefasControl_Giris_FINAL.xlsx"
SHEET_NAME = "Tablo Verisi"
OUTPUT_FILE = "TefasControl_Sonuc_FINAL.xlsx"

# Senin nihai aralığın A6:F558 - kodlar A sütununda
START_ROW = 6
END_ROW = 558

async def check_one(browser, kod):
    url = f"https://www.tefas.gov.tr/tr/fon-detayli-analiz/{kod}"
    page = await browser.new_page(
        user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/122.0.0.0 Safari/537.36"
    )
    try:
        await page.goto(url, wait_until="domcontentloaded", timeout=20000)
        await page.wait_for_timeout(2000)
        content = await page.content()
        if "TEFAS'ta İşlem Görmüyor" in content:
            durum = "Hayır"
        elif "TEFAS'ta İşlem Görüyor" in content:
            durum = "Evet"
        else:
            durum = "Bulunamadı"
        print(f"{kod} -> {durum}")
        await page.close()
        return kod, durum
    except Exception as e:
        print(f"{kod} HATA: {e}")
        try:
            await page.close()
        except:
            pass
        return kod, "HATA"

async def main():
    wb = openpyxl.load_workbook(INPUT_FILE)
    ws = wb[SHEET_NAME]
    
    kodlar = []
    for r in range(START_ROW, END_ROW+1):
        kod = ws.cell(r, 1).value
        if kod:
            kodlar.append(str(kod).strip())
    
    print(f"Toplam {len(kodlar)} fon okundu (A6:A558)")
    
    results = {}
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        for i in range(0, len(kodlar), 5):
            batch = kodlar[i:i+5]
            tasks = [check_one(browser, k) for k in batch]
            batch_results = await asyncio.gather(*tasks)
            for kod, durum in batch_results:
                results[kod] = durum
            await asyncio.sleep(1)
        await browser.close()
    
    # G sütununa yaz - G5 başlık, G6:G558 veri
    ws.cell(5, 7).value = "Platform Durumu"
    for r in range(START_ROW, END_ROW+1):
        kod = ws.cell(r, 1).value
        if kod and str(kod).strip() in results:
            ws.cell(r, 7).value = results[str(kod).strip()]
    
    wb.save(OUTPUT_FILE)
    print(f"Bitti: {OUTPUT_FILE}")
    
    # Özet
    from collections import Counter
    print(Counter(results.values()))

if __name__ == "__main__":
    asyncio.run(main())
